from __future__ import annotations

import csv
from pathlib import Path


def extract_text(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix in {".txt", ".md"}:
        return path.read_text(encoding="utf-8", errors="ignore")
    if suffix == ".csv":
        with path.open("r", encoding="utf-8-sig", errors="ignore", newline="") as handle:
            return "\n".join(" | ".join(row) for row in csv.reader(handle))
    if suffix == ".docx":
        from docx import Document

        document = Document(path)
        return "\n".join(p.text for p in document.paragraphs if p.text.strip())
    if suffix == ".pdf":
        from pypdf import PdfReader

        return "\n\f\n".join(page.extract_text() or "" for page in PdfReader(path).pages)
    if suffix == ".xlsx":
        from openpyxl import load_workbook

        workbook = load_workbook(path, read_only=True, data_only=True)
        lines: list[str] = []
        for sheet in workbook.worksheets:
            lines.append(f"## {sheet.title}")
            lines.extend(
                " | ".join("" if value is None else str(value) for value in row)
                for row in sheet.iter_rows(values_only=True)
            )
        return "\n".join(lines)
    raise ValueError(f"不支持的文件类型: {suffix}")


def split_text(text: str, chunk_size: int = 700, overlap: int = 100) -> list[dict]:
    chunks: list[dict] = []
    section: str | None = None
    pages = text.split("\f")
    has_page_boundaries = len(pages) > 1
    global_offset = 0
    for page_index, page_text in enumerate(pages, 1):
        normalized = "\n".join(line.strip() for line in page_text.splitlines() if line.strip())
        start = 0
        while start < len(normalized):
            end = min(len(normalized), start + chunk_size)
            if end < len(normalized):
                natural_end = max(
                    normalized.rfind("\n", start, end), normalized.rfind("。", start, end)
                )
                if natural_end > start + chunk_size // 2:
                    end = natural_end + 1
            content = normalized[start:end].strip()
            first_line = content.splitlines()[0] if content else ""
            if first_line.startswith("#") or len(first_line) <= 30:
                section = first_line.lstrip("# ") or section
            if content:
                chunks.append(
                    {
                        "content": content,
                        "start_char": global_offset + start,
                        "end_char": global_offset + end,
                        "section": section,
                        "page_number": page_index if has_page_boundaries else None,
                    }
                )
            if end >= len(normalized):
                break
            start = max(start + 1, end - overlap)
        global_offset += len(page_text) + 1
    return chunks
